"""
Import a CCAA DOCX into the project JSON model.

Reads paragraphs and tables in document order, maps cover metadata,
headings and FinanceTable objects. Designed as the seed for a new report.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document as DocxDocument  # type: ignore[import-untyped]

from app.domain.content import ContentBlock, ContentModel
from app.domain.finance import BalanceLine, FinanceTable, PygLine, TableCell, TableColumn
from app.services.balance_from_docx import (
    classify_balance_table,
    extract_balance_lines_from_table,
)
from app.services.pyg_from_docx import (
    classify_pyg_table,
    extract_pyg_lines_from_table,
    infer_income_fact,
)


_HEADING1_STYLES = {"Heading 1"}
_HEADING2_STYLES = {"Heading 2"}
_COVER_STYLES = {"Caption", "Title"}

_NUMERIC_RE = re.compile(
    r"^\s*[\(\-]?\s*[\d]+(?:\.[\d]{3})*(?:,[\d]{1,2})?\s*\)?\s*$"
)
_NOTA_RE = re.compile(r"^(?:Nota\s+)?(\d+)[\.\)]\s+", re.IGNORECASE)
_DATE_LONG_RE = re.compile(
    r"(\d{1,2})\s+de\s+(enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
    r"septiembre|octubre|noviembre|diciembre)\s+de\s+(\d{4})",
    re.IGNORECASE,
)
_MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10, "noviembre": 11,
    "diciembre": 12,
}
_UNIPERSONAL_RE = re.compile(r"sociedad\s+unipersonal|s\.?\s*l\.?\s*u\.?", re.I)
_GESTION_RE = re.compile(r"informe\s+de\s+gesti[oó]n", re.I)
_CUENTAS_RE = re.compile(r"cuentas\s+anuales|ejercicio\s+terminado", re.I)
_BALANCE_TITLE_RE = re.compile(r"balance\s+de\s+situaci[oó]n|^activo$|^pasivo", re.I)


@dataclass
class DocxImportResult:
    content: ContentModel
    tables: dict[str, FinanceTable] = field(default_factory=dict)
    balance_lines: list[BalanceLine] = field(default_factory=list)
    pyg_lines: list[PygLine] = field(default_factory=list)
    facts: dict[str, float] = field(default_factory=dict)
    legal_name: str | None = None
    legal_form_note: str | None = None
    title: str = "Cuentas Anuales"
    detected_current_end: date | None = None
    cover_lines: list[str] = field(default_factory=list)


def _parse_spanish_number(text: str) -> float | None:
    clean = text.strip()
    if not clean or clean in {"-", "–", "—"}:
        return None
    negative = False
    if clean.startswith("(") and clean.endswith(")"):
        negative = True
        clean = clean[1:-1].strip()
    elif clean.startswith("-"):
        negative = True
        clean = clean[1:].strip()
    clean = clean.replace(".", "").replace(",", ".")
    try:
        val = float(clean)
        return -val if negative else val
    except ValueError:
        return None


def _slugify(text: str) -> str:
    slug = text.lower().strip()
    for a, b in (
        ("áàä", "a"), ("éèë", "e"), ("íìï", "i"), ("óòö", "o"), ("úùü", "u"), ("ñ", "n"),
    ):
        for ch in a:
            slug = slug.replace(ch, b)
    slug = re.sub(r"[^a-z0-9]+", "_", slug).strip("_")
    return slug[:50] if slug else "block"


def _make_unique(base_id: str, used: set[str]) -> str:
    candidate = base_id
    counter = 2
    while candidate in used:
        candidate = f"{base_id}_{counter}"
        counter += 1
    used.add(candidate)
    return candidate


def _extract_note_number(text: str, heading_counter: list[int]) -> int:
    m = _NOTA_RE.match(text.strip())
    if m:
        return int(m.group(1))
    heading_counter[0] += 1
    return heading_counter[0]


def _parse_spanish_date(text: str) -> date | None:
    m = _DATE_LONG_RE.search(text)
    if not m:
        return None
    day = int(m.group(1))
    month = _MONTHS[m.group(2).lower()]
    year = int(m.group(3))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def import_docx(docx_path: Path) -> DocxImportResult:
    """Parse a CCAA DOCX into content + tables + cover metadata."""
    doc = DocxDocument(str(docx_path))

    blocks: list[ContentBlock] = []
    tables: dict[str, FinanceTable] = {}
    used_ids: set[str] = set()
    heading_counter = [0]
    current_section = "portada"
    in_cover = True
    cover_lines: list[str] = []
    management_text: str | None = None
    legal_name: str | None = None
    legal_form_note: str | None = None
    detected_end: date | None = None
    table_count = 0
    balance_lines: list[BalanceLine] = []
    pyg_lines: list[PygLine] = []
    facts: dict[str, float] = {}
    seen_balance_sections: set[str] = set()
    pyg_captured = False

    para_map = {id(p._element): p for p in doc.paragraphs}
    table_map = {id(t._element): t for t in doc.tables}

    for child in doc.element.body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

        if tag == "p":
            para = para_map.get(id(child))
            if para is None:
                continue
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else "Normal"
            detected_end = detected_end or _parse_spanish_date(text)

            if in_cover:
                if style_name in _COVER_STYLES or (
                    legal_name is None
                    and len(text) < 120
                    and ("S.L" in text.upper() or "S.A" in text.upper() or text.isupper())
                ):
                    if legal_name is None and not text.startswith("("):
                        legal_name = text
                        continue

                if text.startswith("(") and _UNIPERSONAL_RE.search(text):
                    legal_form_note = text.strip("()")
                    if not legal_form_note.lower().startswith("sociedad"):
                        legal_form_note = text
                    continue

                if _GESTION_RE.search(text):
                    management_text = text
                    continue

                if _CUENTAS_RE.search(text) or cover_lines:
                    # Accumulate subtitle fragments until we leave cover
                    if not _GESTION_RE.search(text):
                        cover_lines.append(text)
                        continue

                # Leaving cover: first real heading or dense body
                if style_name in _HEADING1_STYLES or style_name in _HEADING2_STYLES:
                    in_cover = False
                elif len(text) > 160:
                    in_cover = False
                else:
                    cover_lines.append(text)
                    continue

            if style_name in _HEADING1_STYLES:
                note_num = _extract_note_number(text, heading_counter)
                current_section = f"nota_{note_num}"
                block_id = _make_unique(f"h1.{_slugify(text)}", used_ids)
                # Strip leading "1. " from display if we show note_number separately
                display = _NOTA_RE.sub("", text).strip() or text
                blocks.append(ContentBlock(
                    id=block_id,
                    kind="heading1",
                    text=display,
                    section=current_section,
                    note_number=note_num,
                    keep_with_next=True,
                ))
            elif style_name in _HEADING2_STYLES:
                block_id = _make_unique(f"h2.{_slugify(text)}", used_ids)
                blocks.append(ContentBlock(
                    id=block_id,
                    kind="heading2",
                    text=text,
                    section=current_section,
                    keep_with_next=True,
                ))
            else:
                block_id = _make_unique(f"p.{_slugify(text)}", used_ids)
                blocks.append(ContentBlock(
                    id=block_id,
                    kind="paragraph",
                    text=text,
                    section=current_section or "memoria",
                ))

        elif tag == "tbl":
            tbl_obj = table_map.get(id(child))
            if tbl_obj is None:
                continue
            if in_cover:
                in_cover = False
                current_section = "estados_financieros"

            # Balance tipado: Activo / Pasivo → statements.balance (no tabla genérica)
            balance_section = classify_balance_table(tbl_obj)
            if balance_section and balance_section not in seen_balance_sections:
                extracted = extract_balance_lines_from_table(tbl_obj, balance_section)
                if extracted:
                    balance_lines.extend(extracted)
                    seen_balance_sections.add(balance_section)
                    continue

            # PyG tipado
            if not pyg_captured and classify_pyg_table(tbl_obj):
                extracted_pyg = extract_pyg_lines_from_table(tbl_obj)
                if extracted_pyg:
                    pyg_lines = extracted_pyg
                    pyg_captured = True
                    income = infer_income_fact(extracted_pyg)
                    if income is not None:
                        facts["INGRESOS_N"] = income
                        # N-1 from same line if present
                        for ln in extracted_pyg:
                            if "importe neto" in ln.label.lower() and "cifra" in ln.label.lower():
                                facts["INGRESOS_N1"] = ln.n1
                                break
                    continue

            table_count += 1
            tbl_id = f"table_{table_count}"
            ft = _import_table(tbl_obj, tbl_id, current_section)
            if ft and ft.rows:
                tables[tbl_id] = ft
                ref_id = _make_unique(f"tref.{tbl_id}", used_ids)
                blocks.append(ContentBlock(
                    id=ref_id,
                    kind="table_ref",
                    text=ft.title or f"Tabla {table_count}",
                    section=current_section,
                    table_id=tbl_id,
                ))

    # Semantic cover blocks (stable ids for templates)
    cover_blocks: list[ContentBlock] = []
    if cover_lines:
        subtitle = " ".join(cover_lines)
        subtitle = re.sub(r"\s+", " ", subtitle).strip()
        # Prefer placeholder-friendly wording if we detected a date
        if detected_end and "ejercicio" in subtitle.lower():
            subtitle = (
                "Cuentas Anuales correspondientes al ejercicio terminado el "
                "{{ period.current_end_long }}"
            )
        cover_blocks.append(ContentBlock(
            id="cover.subtitle",
            kind="paragraph",
            text=subtitle,
            section="portada",
            keep_with_next=True,
        ))
        used_ids.add("cover.subtitle")

    if management_text:
        cover_blocks.append(ContentBlock(
            id="cover.management",
            kind="paragraph",
            text=management_text,
            section="portada",
            keep_with_next=True,
        ))
        used_ids.add("cover.management")

    return DocxImportResult(
        content=ContentModel(blocks=cover_blocks + blocks),
        tables=tables,
        balance_lines=balance_lines,
        pyg_lines=pyg_lines,
        facts=facts,
        legal_name=legal_name,
        legal_form_note=legal_form_note,
        detected_current_end=detected_end,
        cover_lines=cover_lines,
    )


_CURRENCY_ONLY_RE = re.compile(r"^\s*euros?\s*$", re.IGNORECASE)
_CURRENCY_PREFIX_RE = re.compile(r"^\s*euros?\s*[—\-–:]\s*", re.IGNORECASE)
_YEAR_RE = re.compile(r"^(19|20)\d{2}$")
_GENERIC_HEADER_RE = re.compile(
    r"^(miles\s+de\s+euros|euros|nota)$",
    re.IGNORECASE,
)
_TEXT_HEADER_HINTS = re.compile(
    r"nombre|direcci[oó]n|actividad|descripci[oó]n|concepto|sociedad",
    re.IGNORECASE,
)


def _clean_header_text(text: str) -> str:
    clean = text.strip()
    if not clean:
        return ""
    if _CURRENCY_ONLY_RE.match(clean):
        return ""
    clean = _CURRENCY_PREFIX_RE.sub("", clean).strip()
    return clean


def _is_year_label(text: str) -> bool:
    return bool(_YEAR_RE.match(text.strip()))


def _is_amount_cell(text: str) -> bool:
    """True for monetary values; bare years are labels/headers."""
    t = text.strip()
    if not t or t in {"-", "–", "—"}:
        return False
    if _is_year_label(t):
        return False
    return _parse_spanish_number(t) is not None


def _row_has_amounts(cells: list[str]) -> bool:
    return any(_is_amount_cell(c) for c in cells[1:])


def _row_is_blank(cells: list[str]) -> bool:
    return not any(c.strip() for c in cells)


def _compose_header(parts: list[str]) -> str:
    """Prefer specific labels over generic 'Miles de euros' bands."""
    cleaned: list[str] = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        if _GENERIC_HEADER_RE.match(p):
            continue
        if p not in cleaned:
            cleaned.append(p)
    if cleaned:
        return cleaned[-1]
    for p in reversed(parts):
        if p.strip():
            return p.strip()
    return ""


def _short_header(text: str, max_len: int = 48) -> str:
    if len(text) <= max_len:
        return text
    return text[: max_len - 1].rstrip() + "…"


def _import_table(
    tbl: object,
    table_id: str,
    section: str,
) -> FinanceTable | None:
    raw_rows = list(tbl.rows)  # type: ignore[attr-defined]
    if len(raw_rows) < 2:
        return None

    matrix: list[tuple[list[str], list[bool]]] = []
    for row in raw_rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        if _row_is_blank(cells):
            continue
        bolds: list[bool] = []
        for cell in row.cells:
            bold = False
            try:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.bold is True:
                            bold = True
                            break
            except Exception:
                pass
            bolds.append(bold)
        while len(bolds) < len(cells):
            bolds.append(False)
        matrix.append((cells, bolds))

    if len(matrix) < 2:
        return None

    num_cols = max(len(cells) for cells, _ in matrix)
    norm: list[tuple[list[str], list[bool]]] = []
    for cells, bolds in matrix:
        cells = cells + [""] * (num_cols - len(cells))
        bolds = bolds + [False] * (num_cols - len(bolds))
        norm.append((cells[:num_cols], bolds[:num_cols]))
    matrix = norm

    header_count = 1
    while header_count < min(4, len(matrix) - 1):
        if _row_has_amounts(matrix[header_count][0]):
            break
        header_count += 1

    header_rows = matrix[:header_count]
    data_matrix = matrix[header_count:]
    if not data_matrix:
        return None

    raw_headers: list[str] = []
    for col_idx in range(num_cols):
        parts = []
        for cells, _ in header_rows:
            part = _clean_header_text(cells[col_idx])
            if part:
                parts.append(part)
        deepest = _clean_header_text(header_rows[-1][0][col_idx])
        if header_count >= 3 and not deepest:
            # Empty cell in the most specific header tier → spacer column
            raw_headers.append("")
        elif deepest:
            raw_headers.append(_short_header(deepest))
        else:
            raw_headers.append(_short_header(_compose_header(parts)))

    keep_idxs: list[int] = []
    for col_idx in range(num_cols):
        header = raw_headers[col_idx]
        has_values = any(
            cells[col_idx].strip() not in {"", "-", "–", "—"}
            for cells, _ in data_matrix
        )
        # Keep labeled columns even if currently empty (maturity schedules, etc.)
        if col_idx == 0 or has_values or header:
            keep_idxs.append(col_idx)

    if not keep_idxs:
        return None

    columns: list[TableColumn] = []
    for new_i, col_idx in enumerate(keep_idxs):
        header = raw_headers[col_idx]
        if not header:
            header = _short_header(_compose_header([
                _clean_header_text(cells[col_idx]) for cells, _ in header_rows
            ]))
        else:
            header = _short_header(header)
        numerics_count = sum(
            1
            for cells, _ in data_matrix
            if _is_amount_cell(cells[col_idx]) or cells[col_idx].strip() in {"-", "–", "—"}
        )
        is_amount = new_i > 0 and numerics_count > max(1, len(data_matrix) // 4)
        is_text = bool(header and _TEXT_HEADER_HINTS.search(header)) and not is_amount
        if new_i == 0:
            align: str = "left"
        elif header.lower() == "nota" or (not is_amount and not is_text and numerics_count == 0):
            align = "center"
        elif is_text:
            align = "left"
        else:
            align = "right"

        columns.append(
            TableColumn(
                id=f"c{new_i}",
                header=header,
                align=align,  # type: ignore[arg-type]
                is_amount=is_amount,
            )
        )

    data_rows: list[dict[str, TableCell]] = []
    total_indices: list[int] = []
    for cells, bolds in data_matrix:
        cells_dict: dict[str, TableCell] = {}
        is_total_row = False
        for new_i, col_idx in enumerate(keep_idxs):
            col_def = columns[new_i]
            cell_text = cells[col_idx]
            bold = bolds[col_idx] if col_idx < len(bolds) else False
            numeric = None
            if col_def.is_amount:
                if cell_text.strip() in {"-", "–", "—"}:
                    numeric = 0.0
                else:
                    numeric = _parse_spanish_number(cell_text)
            is_total = new_i == 0 and cell_text.lower().startswith("total")
            if is_total:
                is_total_row = True
            cells_dict[col_def.id] = TableCell(
                value=cell_text,
                numeric=numeric,
                bold=bold or is_total,
                is_total=is_total,
            )
        data_rows.append(cells_dict)
        if is_total_row:
            total_indices.append(len(data_rows) - 1)

    title = ""
    for row in data_rows:
        cell = row.get("c0")
        if cell and cell.value and len(cell.value.strip()) > 2:
            title = cell.value.strip()[:80]
            break
    if not title:
        for c in columns:
            if c.header and c.header.lower() != "nota":
                title = c.header[:80]
                break

    table_section = section
    if _BALANCE_TITLE_RE.search(title) or any(
        _BALANCE_TITLE_RE.search(c.header) for c in columns[:2]
    ):
        table_section = "estados_financieros"

    return FinanceTable(
        id=table_id,
        title=title,
        columns=columns,
        rows=data_rows,
        section=table_section,
        total_row_indices=total_indices,
    )
