"""
Export a project report to Word (.docx) — estilo CCAA HH Print.

Portada aireada, Arial Narrow, tablas sin grilla pesada, columna Nota,
fechas DD.MM.AAAA, dobles líneas en totales.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document  # type: ignore[import-untyped]
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_COLOR_INDEX  # type: ignore[import-untyped]
from docx.oxml import OxmlElement  # type: ignore[import-untyped]
from docx.oxml.ns import qn  # type: ignore[import-untyped]
from docx.shared import Cm, Pt  # type: ignore[import-untyped]

from app.domain.content import ContentBlock, ContentModel
from app.domain.finance import BalanceLine, FinanceModel, FinanceTable
from app.domain.project import Project
from app.services.formatting import format_amount, format_date_long, format_date_short, format_note_refs

_FONT = "Arial Narrow"


def export_docx(
    project: Project,
    finance: FinanceModel,
    content: ContentModel,
    output_path: Path,
) -> Path:
    doc = Document()
    _set_page(doc)
    _set_default_style(doc)

    _write_cover(doc, project)
    _write_balance(doc, project, finance)
    _write_pyg(doc, project, finance)
    _write_content(doc, finance, content)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def _set_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.3)


def _set_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = _FONT
    style.font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), _FONT)
    rfonts.set(qn("w:hAnsi"), _FONT)
    rfonts.set(qn("w:cs"), "Arial")
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE


def _run(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    size: float = 10,
    highlight: bool = False,
) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = _FONT
    run.font.size = Pt(size)
    if highlight:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), _FONT)
    rfonts.set(qn("w:hAnsi"), _FONT)


def _write_cover(doc: Document, project: Project) -> None:
    # Aire superior tipo portada CCAA
    for _ in range(8):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(title, project.entity.legal_name.upper(), bold=True, size=15)

    if project.entity.legal_form_note:
        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(note, f"({project.entity.legal_form_note})", italic=True, size=11)

    for _ in range(2):
        doc.add_paragraph()

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(sub, project.title, size=12)

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(line, "correspondientes al ejercicio terminado el", size=11)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(date_p, format_date_long(project.period.current_end), bold=True, size=11)

    comp = doc.add_paragraph()
    comp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(
        comp,
        f"(con cifras comparativas al {format_date_long(project.period.prior_end)})",
        italic=True,
        size=9,
    )

    cur = doc.add_paragraph()
    cur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(cur, f"Cifras expresadas en {project.period.currency_label}", size=9)

    doc.add_page_break()


def _set_cell_borders(cell, *, bottom: str | None = "single", bottom_sz: str = "4", top: str | None = None, top_sz: str = "4") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge == "bottom" and bottom:
            el.set(qn("w:val"), bottom)
            el.set(qn("w:sz"), bottom_sz)
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "222222")
        elif edge == "top" and top:
            el.set(qn("w:val"), top)
            el.set(qn("w:sz"), top_sz)
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "222222")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    # replace existing
    existing = tc_pr.find(qn("w:tcBorders"))
    if existing is not None:
        tc_pr.remove(existing)
    tc_pr.append(borders)


def _set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9, align: str = "left") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    _run(p, text, bold=bold, size=size)


def _add_statement_table(
    doc: Document,
    headers: list[str],
    rows: list[tuple[str, str, str, str, bool]],
) -> None:
    """rows: (label, note, n, n1, is_emphasis)"""
    table = doc.add_table(rows=1, cols=4)
    table.autofit = True

    hdr = table.rows[0].cells
    aligns = ["left", "center", "right", "right"]
    for i, h in enumerate(headers):
        _set_cell_text(hdr[i], h, bold=True, size=9, align=aligns[i])
        _set_cell_borders(hdr[i], bottom="single", bottom_sz="8")

    for label, note, n, n1, emph in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], label, bold=emph, size=9, align="left")
        _set_cell_text(cells[1], note, bold=False, size=8, align="center")
        _set_cell_text(cells[2], n, bold=emph, size=9, align="right")
        _set_cell_text(cells[3], n1, bold=emph, size=9, align="right")
        for c in cells:
            _set_cell_borders(c, bottom=None)
        # totals get top+double bottom
        if emph and label.upper().startswith("TOTAL"):
            for c in cells:
                _set_cell_borders(c, top="single", top_sz="6", bottom="double", bottom_sz="6")
        elif emph:
            for c in cells:
                _set_cell_borders(c, top="single", top_sz="4", bottom=None)


def _write_balance(doc: Document, project: Project, finance: FinanceModel) -> None:
    lines = finance.statements.balance.lines
    if not lines:
        return

    head = doc.add_paragraph()
    _run(head, "Balance de situación", bold=True, size=12)
    cap = doc.add_paragraph()
    _run(
        cap,
        f"Al {format_date_long(project.period.current_end)} "
        f"(comparativo {format_date_long(project.period.prior_end)}) — {project.period.currency_label}",
        italic=True,
        size=9,
    )

    d_n = format_date_short(project.period.current_end)
    d_n1 = format_date_short(project.period.prior_end)
    headers = ["", "Nota", d_n, d_n1]

    for section, title in (
        ("asset", "ACTIVO"),
        ("equity_liability", "PATRIMONIO NETO Y PASIVO"),
    ):
        sect = [ln for ln in lines if ln.section == section]
        if not sect:
            continue
        sub = doc.add_paragraph()
        _run(sub, title, bold=True, size=10)
        rows: list[tuple[str, str, str, str, bool]] = []
        for ln in sect:
            if not ln.is_publishable():
                continue
            indent = "    " * max(ln.level - 1, 0)
            emph = ln.role in {"total", "subtotal", "subtotal_group", "header"}
            rows.append(
                (
                    f"{indent}{ln.label}",
                    format_note_refs(ln.note_refs),
                    format_amount(ln.n),
                    format_amount(ln.n1),
                    emph,
                )
            )
        _add_statement_table(doc, headers, rows)
        doc.add_paragraph()

    doc.add_page_break()


def _write_pyg(doc: Document, project: Project, finance: FinanceModel) -> None:
    lines = finance.statements.pyg.lines
    if not lines:
        return

    head = doc.add_paragraph()
    _run(head, "Cuenta de pérdidas y ganancias", bold=True, size=12)
    cap = doc.add_paragraph()
    _run(
        cap,
        f"Ejercicio terminado el {format_date_long(project.period.current_end)} — {project.period.currency_label}",
        italic=True,
        size=9,
    )

    y_n = f"Ejercicio {project.period.current_end.year}"
    y_n1 = f"Ejercicio {project.period.prior_end.year}"
    rows: list[tuple[str, str, str, str, bool]] = []
    for ln in lines:
        if not ln.is_publishable():
            continue
        indent = "    " * max(ln.level - 1, 0)
        emph = ln.role in {"total", "subtotal", "subtotal_group", "header"}
        rows.append(
            (
                f"{indent}{ln.label}",
                format_note_refs(ln.note_refs),
                format_amount(ln.n),
                format_amount(ln.n1),
                emph,
            )
        )
    _add_statement_table(doc, ["", "Nota", y_n, y_n1], rows)
    doc.add_page_break()


def _write_content(doc: Document, finance: FinanceModel, content: ContentModel) -> None:
    skip_ids = {"cover.subtitle", "cover.management", "balance.caption"}
    for block in content.visible_blocks():
        if block.id in skip_ids:
            continue
        _write_block(doc, block, finance)


def _write_block(doc: Document, block: ContentBlock, finance: FinanceModel) -> None:
    if block.kind == "heading1":
        text = block.text
        if block.note_number is not None:
            text = f"{block.note_number}. {text}"
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        _run(p, text, bold=True, size=11)
        return

    if block.kind == "heading2":
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        _run(p, block.text, bold=True, size=10)
        return

    if block.kind == "title":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, block.text, bold=True, size=12)
        return

    if block.kind == "list_item":
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        _run(p, f"• {block.text}", size=10, highlight=block.highlight)
        return

    if block.kind == "table_ref":
        tid = block.table_id
        if tid and tid in finance.tables:
            _write_finance_table(doc, finance.tables[tid])
        return

    if block.text.strip():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _run(p, block.text, size=10, highlight=block.highlight)


def _write_finance_table(doc: Document, tbl: FinanceTable) -> None:
    if tbl.title:
        p = doc.add_paragraph()
        _run(p, tbl.title, bold=True, italic=True, size=9)

    cols = tbl.columns
    if not cols:
        return

    table = doc.add_table(rows=1, cols=len(cols))
    hdr = table.rows[0].cells
    for i, col in enumerate(cols):
        align = "right" if col.is_amount else "left"
        _set_cell_text(hdr[i], col.header, bold=True, size=8, align=align)
        _set_cell_borders(hdr[i], bottom="single", bottom_sz="6")

    for row_idx, row in enumerate(tbl.rows):
        cells = table.add_row().cells
        is_total = row_idx in tbl.total_row_indices
        for i, col in enumerate(cols):
            cell = row.get(col.id)
            align = "right" if col.is_amount else "left"
            if cell is None:
                text = ""
                bold = is_total
            elif col.is_amount and cell.numeric is not None:
                text = format_amount(cell.numeric)
                bold = is_total or cell.bold or cell.is_total
            else:
                text = cell.value
                bold = is_total or cell.bold or cell.is_total
            _set_cell_text(cells[i], text, bold=bold, size=8, align=align)
            if is_total:
                _set_cell_borders(cells[i], top="single", top_sz="4", bottom="double", bottom_sz="4")
            else:
                _set_cell_borders(cells[i], bottom=None)
