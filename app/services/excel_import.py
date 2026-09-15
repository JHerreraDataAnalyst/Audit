from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.domain.finance import BalanceLine, FinanceModel


REQUIRED_HEADERS = ["id", "label", "level", "section", "role", "note_refs", "n", "n1"]


def import_balance_from_excel(path: Path, finance: FinanceModel | None = None) -> FinanceModel:
    """Import balance lines from a simple Excel sheet.

    Expected columns: id, label, level, section, role, note_refs, n, n1
    note_refs may be semicolon-separated (e.g. "8;9;19").
    """
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("El Excel está vacío.")

    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    if missing:
        raise ValueError(f"Faltan columnas: {', '.join(missing)}")

    idx = {h: headers.index(h) for h in REQUIRED_HEADERS}
    lines: list[BalanceLine] = []
    for raw in rows[1:]:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        note_raw = raw[idx["note_refs"]]
        notes: list[str] = []
        if note_raw is not None and str(note_raw).strip():
            notes = [p.strip() for p in str(note_raw).replace(",", ";").split(";") if p.strip()]

        lines.append(
            BalanceLine(
                id=str(raw[idx["id"]]).strip(),
                label=str(raw[idx["label"]]).strip(),
                level=int(raw[idx["level"]] or 1),
                section=str(raw[idx["section"]] or "asset").strip(),  # type: ignore[arg-type]
                role=str(raw[idx["role"]] or "detail").strip(),  # type: ignore[arg-type]
                note_refs=notes,
                n=float(raw[idx["n"]] or 0),
                n1=float(raw[idx["n1"]] or 0),
            )
        )

    model = finance or FinanceModel()
    model.statements.balance.lines = lines
    return model


def write_sample_excel(path: Path) -> Path:
    """Create a sample Excel matching the import schema."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance"
    ws.append(REQUIRED_HEADERS)
    sample = [
        ["intangible", "Inmovilizado intangible", 1, "asset", "subtotal_group", "5", 0, 18106],
        ["software", "Aplicaciones informáticas", 2, "asset", "detail", "", 0, 18106],
        ["ppe", "Inmovilizado material", 1, "asset", "subtotal_group", "6", 105676.29, 13004.78],
        [
            "ppe_other",
            "Instalaciones técnicas y otro inmovilizado material",
            2,
            "asset",
            "detail",
            "",
            105676.29,
            13004.78,
        ],
        ["total_assets", "TOTAL ACTIVO", 1, "asset", "total", "", 4595929.99, 4220600.46],
        ["equity_total", "Patrimonio neto", 1, "equity_liability", "subtotal", "12", 3210450.22, 2980120.15],
        [
            "total_equity_liability",
            "TOTAL PATRIMONIO NETO Y PASIVO",
            1,
            "equity_liability",
            "total",
            "",
            4595929.99,
            4220600.46,
        ],
    ]
    for row in sample:
        ws.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
