"""
Professional Excel export for financial reports.

Generates a multi-sheet Excel workbook with:
1. Control de Cuadre (Audit summary & validation dashboard with hyperlinks)
2. Balance de Situación (Full balance sheet)
3. Cuenta de Pérdidas y Ganancias (PyG)
4. Estado de Cambios en el Patrimonio Neto (ECPN)
5. Estado de Flujos de Efectivo (EFE)
6. All individual notes and detail tables from the Memoria with human-readable sheet names
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from app.domain.finance import FinanceModel, FinanceTable
from app.domain.project import Project
from app.domain.validation import validate_all


# ── Style definitions ─────────────────────────────────────────────────────────

_BRAND_DARK = "1F4B3A"       # Forest green
_BRAND_LIGHT = "E8F0ED"      # Soft sage
_CARD_BG = "F7FAF8"          # Ultra-soft background
_BORDER_COLOR = "D0DCD5"

_HEADER_FILL = PatternFill(start_color=_BRAND_DARK, end_color=_BRAND_DARK, fill_type="solid")
_SUBTOTAL_FILL = PatternFill(start_color="F2F4F3", end_color="F2F4F3", fill_type="solid")
_TOTAL_FILL = PatternFill(start_color=_BRAND_LIGHT, end_color=_BRAND_LIGHT, fill_type="solid")
_OK_FILL = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
_WARN_FILL = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")

_HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
_TITLE_FONT = Font(name="Segoe UI", size=15, bold=True, color=_BRAND_DARK)
_SUBTITLE_FONT = Font(name="Segoe UI", size=11, bold=True, color=_BRAND_DARK)
_SECTION_FONT = Font(name="Segoe UI", size=11, bold=True, color=_BRAND_DARK)
_NORMAL_FONT = Font(name="Segoe UI", size=10)
_BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
_TOTAL_FONT = Font(name="Segoe UI", size=10, bold=True, color=_BRAND_DARK)
_MUTED_FONT = Font(name="Segoe UI", size=9, color="666666")
_LINK_FONT = Font(name="Segoe UI", size=10, color="1565C0", underline="single")
_OK_FONT = Font(name="Segoe UI", size=10, bold=True, color="137333")
_WARN_FONT = Font(name="Segoe UI", size=10, bold=True, color="B06000")

_THIN_BORDER = Border(bottom=Side(style="thin", color="E0E0E0"))
_TOTAL_BORDER = Border(
    top=Side(style="thin", color=_BRAND_DARK),
    bottom=Side(style="double", color=_BRAND_DARK),
)
_HEADER_BORDER = Border(bottom=Side(style="medium", color="FFFFFF"))
_CARD_BORDER = Border(
    top=Side(style="thin", color=_BORDER_COLOR),
    bottom=Side(style="thin", color=_BORDER_COLOR),
    left=Side(style="thin", color=_BORDER_COLOR),
    right=Side(style="thin", color=_BORDER_COLOR),
)

_AMOUNT_NEGATIVE = '#,##0.00;[Red](#,##0.00);"-"'

_RIGHT_ALIGN = Alignment(horizontal="right", vertical="center")
_LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
_CENTER_ALIGN = Alignment(horizontal="center", vertical="center")


# ── Public API ────────────────────────────────────────────────────────────────

def export_excel(
    project: Project,
    finance: FinanceModel,
    output_path: Path,
) -> Path:
    """Generate a professional Excel workbook with all financial statements and notes."""
    wb = Workbook()

    # Identify special statement tables (PyG, Patrimonio, Flujos)
    special_map = _identify_special_tables(finance.tables)

    # 1. Sheet: Balance de Situación
    ws_balance = wb.active
    tot_asset_row, tot_eq_row = _write_balance_sheet(ws_balance, project, finance)

    # 2. Key Statement Sheets (PyG, Patrimonio Neto, Flujos Efectivo)
    created_sheets: list[tuple[str, str, str]] = [("Balance", "Balance de Situación", "Estado")]
    used_sheet_names: set[str] = {"Balance", "Control de Cuadre"}

    for special_name in ["PyG", "Patrimonio Neto", "Flujos Efectivo"]:
        table_id = special_map.get(special_name)
        if table_id and table_id in finance.tables:
            table = finance.tables[table_id]
            clean_name = _safe_sheet_name(special_name, used_sheet_names)
            _write_generic_table_sheet(wb, clean_name, table, project, is_primary_stmt=True)
            created_sheets.append((clean_name, table.title or special_name, "Estado"))

    # 3. Sheets for all remaining tables (Notas de la Memoria)
    special_tids = set(special_map.values())
    for table_id, table in finance.tables.items():
        if table_id in special_tids:
            continue

        clean_title = _derive_clean_title(table, table_id)
        # Prefixed sheet name e.g. "N5 - Intangible" or clean title
        prefix = f"N{_extract_note_number(table.section)} - " if _extract_note_number(table.section) else ""
        raw_name = f"{prefix}{clean_title}"
        sheet_name = _safe_sheet_name(raw_name, used_sheet_names)

        _write_generic_table_sheet(wb, sheet_name, table, project, is_primary_stmt=False)
        created_sheets.append((sheet_name, clean_title, table.section or "Memoria"))

    # 4. Sheet: Control de Cuadre & Resumen Dashboard (inserted at index 0)
    _write_control_dashboard_sheet(wb, project, finance, tot_asset_row, tot_eq_row, created_sheets)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path


# ── Dashboard Sheet ───────────────────────────────────────────────────────────

def _write_control_dashboard_sheet(
    wb: Workbook,
    project: Project,
    finance: FinanceModel,
    tot_asset_row: int,
    tot_eq_row: int,
    sheets: list[tuple[str, str, str]],
) -> None:
    ws = wb.create_sheet(title="Control de Cuadre", index=0)
    ws.views.sheetView[0].showGridLines = True
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 22

    # Title
    ws.merge_cells("B2:E2")
    c = ws.cell(row=2, column=2, value=project.entity.legal_name)
    c.font = _TITLE_FONT

    ws.merge_cells("B3:E3")
    c = ws.cell(row=3, column=2, value=f"CUENTAS ANUALES — EJERCICIO {project.period.current_end}")
    c.font = _SUBTITLE_FONT

    ws.merge_cells("B4:E4")
    c = ws.cell(
        row=4,
        column=2,
        value=f"Control de Cuadre y Resumen Financiero · Cifras en {project.period.currency_label}",
    )
    c.font = _MUTED_FONT

    # Section 1: Cuadre del Balance
    row = 6
    ws.cell(row=row, column=2, value="1. VERIFICACIÓN DE CUADRE DEL BALANCE").font = _SECTION_FONT
    row += 1

    headers = ["Concepto", f"Ejercicio {project.period.current_end}", f"Ejercicio {project.period.prior_end}", "Estado"]
    for ci, h in enumerate(headers, 2):
        cell = ws.cell(row=row, column=ci, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _RIGHT_ALIGN if ci in (3, 4) else _LEFT_ALIGN
    row += 1

    # Row: Total Activo
    r_act = row
    ws.cell(row=row, column=2, value="Total Activo").font = _BOLD_FONT
    c = ws.cell(row=row, column=3, value=f"=Balance!C{tot_asset_row}")
    c.number_format = _AMOUNT_NEGATIVE
    c.font = _BOLD_FONT
    c.alignment = _RIGHT_ALIGN
    c = ws.cell(row=row, column=4, value=f"=Balance!D{tot_asset_row}")
    c.number_format = _AMOUNT_NEGATIVE
    c.font = _BOLD_FONT
    c.alignment = _RIGHT_ALIGN
    ws.cell(row=row, column=5, value="").font = _NORMAL_FONT
    row += 1

    # Row: Total Patrimonio Neto y Pasivo
    r_eq = row
    ws.cell(row=row, column=2, value="Total Patrimonio Neto y Pasivo").font = _BOLD_FONT
    c = ws.cell(row=row, column=3, value=f"=Balance!C{tot_eq_row}")
    c.number_format = _AMOUNT_NEGATIVE
    c.font = _BOLD_FONT
    c.alignment = _RIGHT_ALIGN
    c = ws.cell(row=row, column=4, value=f"=Balance!D{tot_eq_row}")
    c.number_format = _AMOUNT_NEGATIVE
    c.font = _BOLD_FONT
    c.alignment = _RIGHT_ALIGN
    ws.cell(row=row, column=5, value="").font = _NORMAL_FONT
    row += 1

    # Row: Diferencia / Descuadre
    diff_val = round(finance.total_assets_n() - finance.total_equity_liability_n(), 2)
    ws.cell(row=row, column=2, value="DIFERENCIA (Activo - Pasivo/PN)").font = _TOTAL_FONT
    c = ws.cell(row=row, column=3, value=f"=C{r_act}-C{r_eq}")
    c.number_format = _AMOUNT_NEGATIVE
    c.font = _TOTAL_FONT
    c.fill = _TOTAL_FILL
    c.border = _TOTAL_BORDER
    c.alignment = _RIGHT_ALIGN

    c = ws.cell(row=row, column=4, value=f"=D{r_act}-D{r_eq}")
    c.number_format = _AMOUNT_NEGATIVE
    c.font = _TOTAL_FONT
    c.fill = _TOTAL_FILL
    c.border = _TOTAL_BORDER
    c.alignment = _RIGHT_ALIGN

    status_cell = ws.cell(
        row=row,
        column=5,
        value="✓ CUADRADO" if abs(diff_val) < 0.01 else f"⛔ DESCUADRE: {diff_val:,.2f} €",
    )
    status_cell.font = _OK_FONT if abs(diff_val) < 0.01 else _WARN_FONT
    status_cell.fill = _OK_FILL if abs(diff_val) < 0.01 else _WARN_FILL
    status_cell.border = _TOTAL_BORDER
    status_cell.alignment = _CENTER_ALIGN
    row += 3

    # Section 2: Validaciones y Alertas
    validation = validate_all(finance)
    ws.cell(row=row, column=2, value="2. RESUMEN DE AUDITORÍA Y VALIDACIONES").font = _SECTION_FONT
    row += 1

    err_count = sum(1 for i in validation.issues if i.severity == "error")
    warn_count = sum(1 for i in validation.issues if i.severity == "warning")

    kpi_text = f"Estado: {'✓ Todo cuadrado' if err_count == 0 else f'⛔ {err_count} descuadre(s)'}"
    if warn_count > 0:
        kpi_text += f" · {warn_count} aviso(s) para revisión"

    c = ws.cell(row=row, column=2, value=kpi_text)
    c.font = _OK_FONT if err_count == 0 else _WARN_FONT
    c.fill = _OK_FILL if err_count == 0 else _WARN_FILL
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    row += 3

    # Section 3: Índice interactivo de hojas
    ws.cell(row=row, column=2, value="3. ÍNDICE DE ESTADOS Y TABLAS (Clic para navegar)").font = _SECTION_FONT
    row += 1

    idx_headers = ["Hoja Excel", "Descripción / Contenido", "Tipo", "Enlace directo"]
    for ci, h in enumerate(idx_headers, 2):
        cell = ws.cell(row=row, column=ci, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _LEFT_ALIGN
    row += 1

    for sheet_name, desc, kind in sheets:
        ws.cell(row=row, column=2, value=sheet_name).font = _BOLD_FONT
        ws.cell(row=row, column=3, value=desc[:50]).font = _NORMAL_FONT
        ws.cell(row=row, column=4, value=kind).font = _MUTED_FONT

        link_cell = ws.cell(row=row, column=5, value=f"Abrir {sheet_name} →")
        link_cell.hyperlink = f"#'{sheet_name}'!A1"
        link_cell.font = _LINK_FONT
        link_cell.alignment = _LEFT_ALIGN

        for c_idx in (2, 3, 4, 5):
            ws.cell(row=row, column=c_idx).border = _THIN_BORDER
        row += 1


# ── Balance Sheet ─────────────────────────────────────────────────────────────

def _write_balance_sheet(ws, project: Project, finance: FinanceModel) -> tuple[int, int]:
    """Writes the Balance sheet and returns (total_assets_row, total_equity_liability_row)."""
    ws.title = "Balance"
    ws.views.sheetView[0].showGridLines = True

    # Page setup
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:5"

    # Widths
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    # Title block
    row = 1
    ws.merge_cells("A1:D1")
    c = ws.cell(row=row, column=1, value=project.entity.legal_name)
    c.font = _TITLE_FONT
    row += 1

    ws.merge_cells("A2:D2")
    c = ws.cell(row=row, column=1, value=f"Balance de Situación al {project.period.current_end}")
    c.font = _SUBTITLE_FONT
    row += 1

    ws.merge_cells("A3:D3")
    c = ws.cell(row=row, column=1, value=f"(Expresado en {project.period.currency_label})")
    c.font = _MUTED_FONT
    row += 2

    # ACTIVO section
    row, tot_asset_row = _write_balance_section(ws, row, "ACTIVO", finance, "asset", project)
    row += 1

    # PATRIMONIO NETO Y PASIVO section
    row, tot_eq_row = _write_balance_section(ws, row, "PATRIMONIO NETO Y PASIVO", finance, "equity_liability", project)

    ws.freeze_panes = "A6"
    return tot_asset_row, tot_eq_row


def _write_balance_section(
    ws,
    start_row: int,
    title: str,
    finance: FinanceModel,
    section: str,
    project: Project,
) -> tuple[int, int]:
    row = start_row
    tot_row = start_row

    # Section header
    headers = [title, "Nota", str(project.period.current_end), str(project.period.prior_end)]
    for col_idx, header in enumerate(headers, 1):
        c = ws.cell(row=row, column=col_idx, value=header)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.border = _HEADER_BORDER
        c.alignment = _CENTER_ALIGN if col_idx == 2 else (_RIGHT_ALIGN if col_idx > 2 else _LEFT_ALIGN)
    row += 1

    lines = finance.lines_by_section(section)
    for line in lines:
        label = line.label
        if line.level == 2:
            label = f"    {label}"
        elif line.level == 3:
            label = f"        {label}"

        c = ws.cell(row=row, column=1, value=label)
        if line.role in ("total", "subtotal"):
            c.font = _TOTAL_FONT
            c.fill = _TOTAL_FILL
            c.border = _TOTAL_BORDER
            if line.role == "total":
                tot_row = row
        elif line.role == "subtotal_group":
            c.font = _BOLD_FONT
        else:
            c.font = _NORMAL_FONT
            c.border = _THIN_BORDER
        c.alignment = _LEFT_ALIGN

        # Note
        note_val = "; ".join(line.note_refs) if line.note_refs else ""
        c = ws.cell(row=row, column=2, value=note_val)
        c.font = _MUTED_FONT
        c.alignment = _CENTER_ALIGN
        if line.role in ("total", "subtotal"):
            c.fill = _TOTAL_FILL
            c.border = _TOTAL_BORDER

        # Amounts
        for col_idx, amount in [(3, line.n), (4, line.n1)]:
            c = ws.cell(row=row, column=col_idx, value=amount)
            c.number_format = _AMOUNT_NEGATIVE
            c.alignment = _RIGHT_ALIGN
            if line.role in ("total", "subtotal"):
                c.font = _TOTAL_FONT
                c.fill = _TOTAL_FILL
                c.border = _TOTAL_BORDER
            elif line.role == "subtotal_group":
                c.font = _BOLD_FONT
            else:
                c.font = _NORMAL_FONT
                c.border = _THIN_BORDER

        row += 1

    return row, tot_row


# ── Generic Statement / Note Table Sheet ──────────────────────────────────────

def _write_generic_table_sheet(
    wb: Workbook,
    sheet_name: str,
    table: FinanceTable,
    project: Project,
    is_primary_stmt: bool = False,
) -> None:
    ws = wb.create_sheet(title=sheet_name)
    ws.views.sheetView[0].showGridLines = True
    ws.page_setup.orientation = "landscape" if len(table.columns) > 4 else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    # Title
    row = 1
    max_col = max(len(table.columns), 3)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    title_text = table.title or sheet_name
    c = ws.cell(row=row, column=1, value=title_text)
    c.font = _TITLE_FONT if is_primary_stmt else _SUBTITLE_FONT
    row += 1

    subtitle = f"{project.entity.legal_name} · Ejercicio {project.period.current_end}"
    if table.section:
        subtitle += f" · Sección: {table.section.replace('_', ' ').capitalize()}"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
    c = ws.cell(row=row, column=1, value=subtitle)
    c.font = _MUTED_FONT
    row += 2

    header_row = row

    # Column headers
    for col_idx, col in enumerate(table.columns, 1):
        header_val = col.header
        if header_val.startswith("Col ") and col_idx == 1:
            header_val = "Concepto"
        c = ws.cell(row=row, column=col_idx, value=header_val)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.border = _HEADER_BORDER
        c.alignment = _RIGHT_ALIGN if col.is_amount else _LEFT_ALIGN

        # Width calculation
        width = max(12, min(50, len(str(header_val)) + 5))
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    row += 1

    # Ensure column 1 is wide enough for concepts
    if table.columns:
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 48)

    # Data rows
    for row_idx, data_row in enumerate(table.rows):
        is_total_row = row_idx in table.total_row_indices

        for col_idx, col in enumerate(table.columns, 1):
            cell = data_row.get(col.id)
            if not cell:
                continue

            if col.is_amount and cell.numeric is not None:
                c = ws.cell(row=row, column=col_idx, value=cell.numeric)
                c.number_format = _AMOUNT_NEGATIVE
                c.alignment = _RIGHT_ALIGN
            else:
                c = ws.cell(row=row, column=col_idx, value=cell.value)
                c.alignment = _LEFT_ALIGN if col_idx == 1 else _RIGHT_ALIGN

            if is_total_row or cell.is_total or cell.bold:
                c.font = _TOTAL_FONT
                c.fill = _TOTAL_FILL
                c.border = _TOTAL_BORDER
            else:
                c.font = _NORMAL_FONT
                c.border = _THIN_BORDER

        row += 1

    # Freeze header
    ws.freeze_panes = f"A{header_row + 1}"


# ── Title & Statement Detectors ───────────────────────────────────────────────

def _identify_special_tables(tables: dict[str, FinanceTable]) -> dict[str, str]:
    """Map primary financial statement names ('PyG', 'Patrimonio Neto', 'Flujos Efectivo') to table IDs."""
    mapping: dict[str, str] = {}

    for tid, t in tables.items():
        title_lower = (t.title or "").lower()
        first_labels = " ".join(
            r.get("c0", {}).value.lower() for r in t.rows[:5] if r.get("c0")
        )
        col_headers = " ".join(c.header.lower() for c in t.columns)

        combined = f"{title_lower} {first_labels} {col_headers}"

        # PyG
        if "pyg" not in mapping:
            if "operaciones continuadas" in combined or "cifra de negocios" in combined or "pérdidas y ganancias" in combined:
                mapping["PyG"] = tid
                continue

        # Patrimonio Neto
        if "Patrimonio Neto" not in mapping:
            if "patrimonio neto" in combined or "fondos propios" in combined or "saldos al" in combined:
                mapping["Patrimonio Neto"] = tid
                continue

        # Flujos de Efectivo
        if "Flujos Efectivo" not in mapping:
            if "flujos de efectivo" in combined or "resultado del ejercicio antes de impuestos" in combined:
                mapping["Flujos Efectivo"] = tid
                continue

    return mapping


def _derive_clean_title(table: FinanceTable, table_id: str) -> str:
    """Derive a friendly, clean title for a table."""
    if table.title and not table.title.lower().startswith("table_"):
        return table.title.strip()

    # Look for first non-empty label cell
    for r in table.rows:
        for col in table.columns[:2]:
            c = r.get(col.id)
            if c and c.value and len(c.value.strip()) > 2 and not c.value.strip().isdigit():
                return c.value.strip()[:60]

    # Look for first non-generic column header
    for col in table.columns:
        if col.header and not col.header.startswith("Col "):
            return col.header.strip()[:60]

    return table_id


def _extract_note_number(section: str) -> str:
    m = re.search(r"\d+", section or "")
    return m.group(0) if m else ""


def _safe_sheet_name(name: str, used: set[str]) -> str:
    """Create a valid, unique Excel sheet name (max 31 chars, no illegal characters)."""
    clean = re.sub(r"[\[\]:*?/\\]", "", name).strip()
    if not clean:
        clean = "Tabla"
    clean = clean[:28]

    candidate = clean
    counter = 2
    while candidate in used:
        candidate = f"{clean[:25]}_{counter}"
        counter += 1
    used.add(candidate)
    return candidate
