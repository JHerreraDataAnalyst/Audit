"""
Comprehensive test suite for the Audit Report system:
- Financial domain models (Balance, Tables, Validation)
- Cross-reference and math mismatch detection
- Excel export (multi-sheet professional workbook, formulas, sheets)
- Table update API (both generic tables and balance statements)
- HTML report rendering
"""

import json
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from starlette.testclient import TestClient

from app.domain.finance import FinanceModel, FinanceTable, TableCell, TableColumn
from app.domain.validation import find_concept_impact, validate_all, validate_balance, validate_table_totals
from app.main import app
from app.services.excel_export import export_excel
from app.services.render import render_report_html
from app.services.storage import load_content, load_finance, load_project


class TestAuditReportSystem(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)
        cls.project = load_project("demo-hh-print")
        cls.finance = load_finance("demo-hh-print")
        cls.content = load_content("demo-hh-print")

    def test_balance_validation(self):
        """Test that the loaded project balance is valid and balanced."""
        res = validate_balance(self.finance)
        self.assertTrue(res.ok, "Balance should be balanced")
        self.assertAlmostEqual(self.finance.total_assets_n(), self.finance.total_equity_liability_n(), delta=0.01)

    def test_validation_detects_imbalance(self):
        """Test that introducing an artificial imbalance is detected."""
        f_copy = self.finance.model_copy(deep=True)
        # Modify asset total line
        total_line = next(ln for ln in f_copy.statements.balance.lines if ln.section == "asset" and ln.role == "total")
        total_line.n += 50000.0
        res = validate_balance(f_copy)
        self.assertFalse(res.ok)
        error_issues = [i for i in res.issues if i.severity == "error"]
        self.assertTrue(len(error_issues) > 0)
        self.assertEqual(error_issues[0].code, "BALANCE_NOT_BALANCED")

    def test_concept_impact_lookup(self):
        """Test that find_concept_impact finds occurrences across Balance and Tables."""
        impacts = find_concept_impact(self.finance, "inmovilizado")
        self.assertIsInstance(impacts, list)
        self.assertGreater(len(impacts), 0)
        sources = [i["source"] for i in impacts]
        self.assertTrue(any("Balance" in s for s in sources))

    def test_excel_export_structure(self):
        """Test that professional Excel workbook is generated with all required sheets."""
        import openpyxl

        out_path = Path("outputs/test_suite_export.xlsx")
        export_excel(self.project, self.finance, out_path)
        self.assertTrue(out_path.exists())
        self.assertGreater(out_path.stat().st_size, 10000)

        wb = openpyxl.load_workbook(str(out_path), data_only=False)
        sheet_names = wb.sheetnames

        # Key sheets must exist
        self.assertIn("Control de Cuadre", sheet_names)
        self.assertIn("Balance", sheet_names)
        self.assertIn("PyG", sheet_names)
        self.assertIn("Patrimonio Neto", sheet_names)
        self.assertIn("Flujos Efectivo", sheet_names)

        # Total sheets should include all notes
        self.assertGreaterEqual(len(sheet_names), 50)

        # Check Control de Cuadre sheet formulas
        ws_ctrl = wb["Control de Cuadre"]
        val_status = ws_ctrl["E10"].value
        self.assertEqual(val_status, "✓ CUADRADO")

    def test_html_render(self):
        """Test that render_report_html renders HTML with editable attributes and tables."""
        html = render_report_html(self.project, self.finance, self.content, editable=True)
        self.assertIn("data-project-id", html)
        self.assertIn("is-editable-table", html)
        self.assertIn("balance_asset", html)
        self.assertIn("table-modal", html)

    def test_update_balance_via_api(self):
        """Test updating a balance line via the POST /tables/balance_asset endpoint."""
        # Find original value of first asset line
        orig_val = self.finance.statements.balance.lines[0].n
        line_id = self.finance.statements.balance.lines[0].id

        # Update via API
        resp = self.client.post(
            "/projects/demo-hh-print/tables/balance_asset",
            json={"cells": [{"row": 0, "col": "n", "line_id": line_id, "numeric": orig_val, "value": str(orig_val)}]},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data["ok"])
        self.assertIn("validation", data)

    def test_update_generic_table_via_api(self):
        """Test updating a generic table cell via the POST /tables/{id} endpoint."""
        first_tid = next(iter(self.finance.tables.keys()))
        resp = self.client.post(
            f"/projects/demo-hh-print/tables/{first_tid}",
            json={"cells": [{"row": 0, "col": "c1", "numeric": 1234.56, "value": "1.234,56"}]},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data["ok"])

    def test_concept_impact_endpoint(self):
        """Test GET /projects/{id}/concept-impact endpoint."""
        resp = self.client.get("/projects/demo-hh-print/concept-impact?label=intangible")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn("impacts", data)

    def test_rollover_shifts_balance_and_period(self):
        from datetime import date

        from app.services.rollover import rollover

        project = self.project.model_copy(deep=True)
        finance = self.finance.model_copy(deep=True)
        old_current = project.period.current_end
        first = finance.statements.balance.lines[0]
        old_n = first.n

        new_project, new_finance = rollover(project, finance, clear_current=True)
        self.assertEqual(new_project.period.prior_end, old_current)
        self.assertEqual(new_project.period.current_end, date(old_current.year + 1, old_current.month, old_current.day))
        rolled = new_finance.statements.balance.lines[0]
        self.assertEqual(rolled.n1, old_n)
        self.assertEqual(rolled.n, 0.0)

    def test_docx_import_cover_semantics(self):
        from pathlib import Path

        from app.services.docx_importer import import_docx

        docx = next(Path("reference").glob("*.docx"))
        result = import_docx(docx)
        self.assertTrue(result.legal_name)
        self.assertIn("HH PRINT", result.legal_name.upper())
        ids = {b.id for b in result.content.blocks}
        self.assertIn("cover.subtitle", ids)
        self.assertGreater(len(result.tables), 10)
        self.assertGreater(len(result.content.blocks), 50)
        # Company name must not become nota_1 heading
        first_h1 = next(b for b in result.content.blocks if b.kind == "heading1")
        self.assertNotIn("HH PRINT", first_h1.text.upper())

    def test_docx_imports_typed_balance(self):
        from pathlib import Path

        from app.domain.finance import BalanceStatement, FinanceModel, Statements
        from app.domain.validation import validate_balance
        from app.services.docx_importer import import_docx

        docx = next(Path("reference").glob("*.docx"))
        result = import_docx(docx)
        self.assertGreaterEqual(len(result.balance_lines), 20)
        sections = {ln.section for ln in result.balance_lines}
        self.assertIn("asset", sections)
        self.assertIn("equity_liability", sections)
        totals = [ln for ln in result.balance_lines if ln.role == "total"]
        self.assertGreaterEqual(len(totals), 2)

        finance = FinanceModel(
            tables=result.tables,
            statements=Statements(balance=BalanceStatement(lines=result.balance_lines)),
        )
        self.assertAlmostEqual(
            finance.total_assets_n(),
            finance.total_equity_liability_n(),
            delta=0.05,
        )
        self.assertTrue(validate_balance(finance).ok)

    def test_docx_imports_typed_pyg(self):
        from pathlib import Path

        from app.services.docx_importer import import_docx
        from app.services.render import render_report_html
        from app.domain.finance import (
            BalanceStatement,
            FinanceModel,
            PygStatement,
            Statements,
        )
        from app.domain.project import Entity, Period, Project
        from datetime import date

        docx = next(Path("reference").glob("*.docx"))
        result = import_docx(docx)
        self.assertGreaterEqual(len(result.pyg_lines), 15)
        roles = {ln.role for ln in result.pyg_lines}
        self.assertIn("total", roles)
        self.assertIn("subtotal_group", roles)
        self.assertIn("INGRESOS_N", result.facts)
        self.assertAlmostEqual(result.facts["INGRESOS_N"], 26573521.07, delta=0.01)

        resultado = next(ln for ln in result.pyg_lines if ln.role == "total")
        self.assertIn("ejercicio", resultado.label.lower())
        self.assertAlmostEqual(resultado.n, 285998.32, delta=0.01)

        project = Project(
            id="tmp-pyg",
            entity=Entity(legal_name=result.legal_name or "Test"),
            period=Period(current_end=date(2023, 3, 31), prior_end=date(2022, 3, 31)),
        )
        finance = FinanceModel(
            facts=result.facts,
            tables=result.tables,
            statements=Statements(
                balance=BalanceStatement(lines=result.balance_lines),
                pyg=PygStatement(lines=result.pyg_lines),
            ),
        )
        html = render_report_html(project, finance, result.content, editable=False)
        self.assertIn("Cuenta de pérdidas y ganancias", html)
        self.assertIn("pyg_statement", html)
        self.assertIn("Balance de situación", html)


if __name__ == "__main__":
    unittest.main()
